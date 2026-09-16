"""Retrieval pipeline: corpus loading, segmentation, chunking, indexing, search.

No model is needed. A deterministic bag-of-words embedder stands in for
nomic-embed-text, and the Ollama client is exercised against a mock transport.
"""

from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
from typing import Callable, List, Optional, Sequence

import httpx
import pytest

from procurecheck.retrieval import chunking
from procurecheck.retrieval.chunking import chunk_document
from procurecheck.retrieval.corpus import (
    EXTRACTION_NATIVE,
    EXTRACTION_OCR,
    KIND_SYNTHETIC,
    CorpusDocument,
    CorpusError,
    load_corpus,
)
from procurecheck.retrieval.embeddings import (
    DOCUMENT_PREFIX,
    QUERY_PREFIX,
    EmbeddingUnavailableError,
    OllamaEmbedder,
    normalise,
)
from procurecheck.retrieval.index import CorpusIndex, IndexNotFoundError
from procurecheck.retrieval.lexical import BM25Index, tokenize
from procurecheck.retrieval.retriever import MODE_BM25, MODE_HYBRID, Retriever
from procurecheck.retrieval.segment import Block, looks_like_heading, segment

DIMENSIONS = 64


class FakeEmbedder:
    """Hashes each token into a fixed-size vector, so shared words mean similarity."""

    model = "fake-embed"

    def _vector(self, text: str) -> List[float]:
        vector = [0.0] * DIMENSIONS
        for token in tokenize(text):
            bucket = int(hashlib.md5(token.encode()).hexdigest(), 16) % DIMENSIONS
            vector[bucket] += 1.0
        return normalise(vector)

    def embed_documents(
        self, texts: Sequence[str], on_batch: Optional[Callable[[int, int], None]] = None
    ) -> List[List[float]]:
        return [self._vector(t) for t in texts]

    def embed_query(self, text: str) -> List[float]:
        return self._vector(text)


def _write_manifest(corpus: Path, rows: Sequence[tuple]) -> None:
    with (corpus / "CORPUS_MANIFEST.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Document ID", "Document Title", "File Name"])
        writer.writerows(rows)


def _document(doc_id: str, source: Path, title: str = "Test Document") -> CorpusDocument:
    return CorpusDocument(doc_id, title, source, source, EXTRACTION_NATIVE)


@pytest.fixture
def corpus(tmp_path: Path) -> Path:
    documents = tmp_path / "documents"
    documents.mkdir()
    (documents / "DOC-001_Guide.txt").write_text(
        "Bid Security\n\nA bid security shall be a bank guarantee valid for 28 days.\n\n"
        "Eligible Bidders\n\nA bidder shall not be debarred or suspended by the Authority.",
        encoding="utf-8",
    )
    (documents / "DOC-002_Scan.pdf").write_bytes(b"%PDF-1.4 placeholder")
    ocr = tmp_path / "derived" / "ocr" / "DOC-002"
    ocr.mkdir(parents=True)
    (ocr / "DOC-002_p01.txt").write_text(
        "ADVANCE PAYMENT SECURITIES\nAn advance payment security shall be\nprovided by the provider.\n",
        encoding="utf-8",
    )
    (documents / "DOC-003_Synthetic.txt").write_text("A current tax clearance certificate is attached to this bid.", encoding="utf-8")
    _write_manifest(
        tmp_path,
        [
            ("DOC-001", "User Guide", "DOC-001_Guide.txt"),
            ("DOC-002", "Guideline on Securities", "DOC-002_Scan.pdf"),
            ("DOC-003", "Synthetic Tender Submission A", "DOC-003_Synthetic.txt"),
        ],
    )
    with (tmp_path / "DERIVATIONS.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Document ID", "Original file", "Derived text source", "Method", "Tool", "Date", "Note"])
        writer.writerow(["DOC-002", "documents/DOC-002_Scan.pdf", "derived/ocr/DOC-002", "OCR (one text file per page)", "test", "2026-09-17", ""])
    return tmp_path


# Corpus ------------------------------------------------------------------


def test_load_corpus_resolves_derived_ocr_text(corpus: Path) -> None:
    documents = {d.doc_id: d for d in load_corpus(corpus)}

    assert documents["DOC-001"].extraction == EXTRACTION_NATIVE
    assert documents["DOC-002"].extraction == EXTRACTION_OCR
    assert documents["DOC-002"].text_source.is_dir()
    assert documents["DOC-003"].kind == KIND_SYNTHETIC


def test_load_corpus_fails_when_a_listed_file_is_missing(corpus: Path) -> None:
    (corpus / "documents" / "DOC-001_Guide.txt").unlink()

    with pytest.raises(CorpusError, match="DOC-001"):
        load_corpus(corpus)


def test_load_corpus_rejects_an_unknown_document_filter(corpus: Path) -> None:
    with pytest.raises(CorpusError, match="DOC-999"):
        load_corpus(corpus, ["DOC-001", "DOC-999"])


# Segmentation ------------------------------------------------------------


@pytest.mark.parametrize(
    "text, expected",
    [
        ("Eligible Bidders", True),
        ("SECTION 3 EVALUATION", True),
        ("Section 1: Instructions to Bidders", True),
        ("Bankrupt; or", False),
        ("Being wound up", False),
        ("The bidder is not:", False),
        ("Corrupt and Fraudulent Practices", True),
        ("Scope", False),
    ],
)
def test_looks_like_heading(text: str, expected: bool) -> None:
    assert looks_like_heading(text) is expected


def test_docx_segmentation_keeps_body_order_pages_and_headings(tmp_path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_paragraph("Table of Contents\t1", style="Normal")
    document.add_heading("Bid Security", level=2)
    document.add_paragraph("The bidder shall furnish a bid security.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Form B-3 Attached"
    table.rows[0].cells[1].text = "No"
    document.add_paragraph("Page 7 of 75")
    document.add_heading("Eligible Bidders", level=2)
    document.add_paragraph("A bidder shall not be suspended.")
    path = tmp_path / "doc.docx"
    document.save(str(path))

    blocks = segment(_document("DOC-009", path))
    texts = [b.text for b in blocks]

    assert "Table of Contents 1" not in texts
    assert texts.index("Form B-3 Attached | No") == texts.index("The bidder shall furnish a bid security.") + 1
    assert "Page 7 of 75" not in texts
    security = next(b for b in blocks if b.text.startswith("The bidder shall furnish"))
    assert security.page == 7 and security.heading == "Bid Security"
    suspended = next(b for b in blocks if b.text.startswith("A bidder shall not"))
    assert suspended.page is None and suspended.heading == "Eligible Bidders"


def test_pdf_layout_drops_running_headers_and_keeps_numbered_titles(tmp_path: Path) -> None:
    import pymupdf

    document = pymupdf.open()
    for number in range(1, 4):
        page = document.new_page()
        page.insert_text((72, 40), "The Public Procurement Guidelines, 2024")
        page.insert_text((72, 120), f"{number}.")
        page.insert_text((72, 134), f"Rules for method {number}.")
        page.insert_text((72, 148), "(1) A procuring entity shall keep records.")
    path = tmp_path / "rules.pdf"
    document.save(str(path))

    blocks = segment(_document("DOC-030", path))

    assert all("Guidelines, 2024" not in b.text for b in blocks)
    assert [b.page for b in blocks] == [1, 2, 3]
    assert blocks[2].heading == "3. Rules for method 3."


def test_ocr_segmentation_reads_one_file_per_page(corpus: Path) -> None:
    ocr_document = next(d for d in load_corpus(corpus) if d.doc_id == "DOC-002")

    blocks = segment(ocr_document)

    assert [b.page for b in blocks] == [1]
    assert blocks[0].heading == "ADVANCE PAYMENT SECURITIES"


# Chunking ----------------------------------------------------------------


def _blocks(heading: str, sentences: int, page: int = 1) -> List[Block]:
    return [Block(f"Sentence {i} about {heading.lower()} requirements here.", page, heading) for i in range(sentences)]


def test_chunks_never_cross_a_heading() -> None:
    blocks = [*_blocks("Bid Security", 3), *_blocks("Eligible Bidders", 3, page=2)]

    chunks = chunk_document(_document("DOC-001", Path("x")), blocks)

    assert len(chunks) == 2
    assert {c.section_heading for c in chunks} == {"Bid Security", "Eligible Bidders"}
    assert all("security" not in c.text or c.section_heading == "Bid Security" for c in chunks)


def test_chunk_ids_are_ordinal_and_neighbours_are_linked() -> None:
    blocks = [*_blocks("A Heading", 2), *_blocks("B Heading", 2), *_blocks("C Heading", 2)]

    chunks = chunk_document(_document("DOC-004", Path("x")), blocks)

    assert [c.chunk_id for c in chunks] == ["DOC-004-C0001", "DOC-004-C0002", "DOC-004-C0003"]
    assert chunks[0].prev_id is None and chunks[0].next_id == "DOC-004-C0002"
    assert chunks[2].prev_id == "DOC-004-C0002" and chunks[2].next_id is None


def test_oversized_block_is_split_within_the_word_limit() -> None:
    sentence = "The provider shall deliver the supplies to the entity. "
    block = Block(sentence * 120, 3, "Delivery")

    chunks = chunk_document(_document("DOC-005", Path("x")), [block])

    assert len(chunks) > 1
    assert all(c.word_count <= chunking.MAX_WORDS for c in chunks)
    assert sum(c.word_count for c in chunks) == len(block.text.split())


def test_index_text_carries_title_and_heading_but_text_stays_verbatim() -> None:
    blocks = _blocks("Bid Security", 2)

    chunk = chunk_document(_document("DOC-006", Path("x"), "Standard Bidding Document"), blocks)[0]

    assert chunk.index_text.startswith("Standard Bidding Document > Bid Security\n")
    assert chunk.text == "\n".join(b.text for b in blocks)
    assert chunk.location == "DOC-006, p. 1, Bid Security"


# Keyword index -----------------------------------------------------------


def test_bm25_ranks_the_passage_sharing_rare_terms_first() -> None:
    index = BM25Index.build(
        [
            "The bidder shall pay taxes and social security contributions.",
            "A bidder that is debarred or suspended is not eligible.",
            "Bid securities shall be bank guarantees.",
        ]
    )

    scores = index.score("debarred bidder")
    assert max(scores, key=scores.get) == 1
    assert 2 in index.score("bid security guarantee")


def test_tokenize_folds_plurals_and_drops_stopwords() -> None:
    assert tokenize("The Securities and the Guarantees of bids") == ["security", "guarantee", "bid"]


# Index and search --------------------------------------------------------


@pytest.fixture
def built_index(corpus: Path) -> CorpusIndex:
    return CorpusIndex.build(load_corpus(corpus), FakeEmbedder())


def test_index_round_trips_through_disk(built_index: CorpusIndex, tmp_path: Path) -> None:
    built_index.save(tmp_path / "index")

    loaded = CorpusIndex.load(tmp_path / "index")

    assert [c.chunk_id for c in loaded.chunks] == [c.chunk_id for c in built_index.chunks]
    assert loaded.has_vectors
    assert loaded.vectors[0] == pytest.approx(built_index.vectors[0], abs=1e-6)
    assert loaded.bm25.score("bank guarantee") == pytest.approx(built_index.bm25.score("bank guarantee"))
    assert json.loads((tmp_path / "index" / "meta.json").read_text())["embedding_model"] == "fake-embed"


def test_loading_a_missing_index_says_how_to_build_it(tmp_path: Path) -> None:
    with pytest.raises(IndexNotFoundError, match="run.py index"):
        CorpusIndex.load(tmp_path / "nothing")


def test_index_detects_a_changed_corpus(built_index: CorpusIndex, corpus: Path) -> None:
    assert not built_index.is_stale(load_corpus(corpus))

    (corpus / "documents" / "DOC-001_Guide.txt").write_text("Changed.", encoding="utf-8")

    assert built_index.is_stale(load_corpus(corpus))


def test_hybrid_search_returns_sources_and_a_trace(built_index: CorpusIndex) -> None:
    result = Retriever(built_index, FakeEmbedder(), min_cosine=0.2).search("debarred or suspended bidder", top_k=2)

    top = result.hits[0]
    assert result.mode == MODE_HYBRID
    assert top.chunk.doc_id == "DOC-001" and "debarred" in top.chunk.text
    assert top.bm25_rank == 1 and top.dense_rank is not None
    assert result.evidence_sufficient is True
    trace = result.to_trace()
    assert trace["hits"][0]["chunk_id"] == top.chunk.chunk_id
    assert trace["hits"][0]["location"].startswith("DOC-001")


def test_search_falls_back_to_keywords_without_an_embedder(built_index: CorpusIndex) -> None:
    result = Retriever(built_index, embedder=None).search("bank guarantee")

    assert result.mode == MODE_BM25
    assert any("skipped" in note for note in result.notes)
    assert result.evidence_sufficient is None
    assert result.hits[0].dense_rank is None


def test_unrelated_query_is_flagged_as_insufficient_evidence(built_index: CorpusIndex) -> None:
    result = Retriever(built_index, FakeEmbedder(), min_cosine=0.5).search("football world cup winners")

    assert result.evidence_sufficient is False


def test_document_filter_limits_results(built_index: CorpusIndex) -> None:
    result = Retriever(built_index, FakeEmbedder()).search("security", doc_ids=["DOC-002"])

    assert result.hits and {h.chunk.doc_id for h in result.hits} == {"DOC-002"}
    assert result.hits[0].chunk.extraction == EXTRACTION_OCR


def test_neighbours_widen_a_hit_in_document_order(built_index: CorpusIndex) -> None:
    first, second = [c for c in built_index.chunks if c.doc_id == "DOC-001"]

    window = Retriever(built_index).neighbours(second)

    assert window == [first, second]


def test_empty_query_is_rejected(built_index: CorpusIndex) -> None:
    with pytest.raises(ValueError):
        Retriever(built_index).search("   ")


# Embedding client --------------------------------------------------------


def _client(handler: Callable[[httpx.Request], httpx.Response]) -> httpx.Client:
    return httpx.Client(transport=httpx.MockTransport(handler))


def test_embedder_applies_task_prefixes_and_normalises() -> None:
    seen: List[dict] = []

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        seen.append(body)
        return httpx.Response(200, json={"embeddings": [[3.0, 4.0] for _ in body["input"]]})

    embedder = OllamaEmbedder("http://ollama", "nomic-embed-text", client=_client(handler))

    documents = embedder.embed_documents(["passage one", "passage two"])
    query = embedder.embed_query("question")

    assert seen[0]["input"] == [DOCUMENT_PREFIX + "passage one", DOCUMENT_PREFIX + "passage two"]
    assert seen[1]["input"] == [QUERY_PREFIX + "question"]
    assert documents[0] == pytest.approx([0.6, 0.8])
    assert query == pytest.approx([0.6, 0.8])


def test_embedder_reports_a_missing_model_with_the_pull_command() -> None:
    embedder = OllamaEmbedder(
        "http://ollama", "nomic-embed-text", client=_client(lambda r: httpx.Response(404, json={}))
    )

    with pytest.raises(EmbeddingUnavailableError, match="ollama pull nomic-embed-text"):
        embedder.embed_query("anything")


def test_embedder_rejects_a_short_response() -> None:
    embedder = OllamaEmbedder(
        "http://ollama", "m", client=_client(lambda r: httpx.Response(200, json={"embeddings": []}))
    )

    with pytest.raises(EmbeddingUnavailableError, match="received 0"):
        embedder.embed_documents(["one"])


def test_rebuild_reuses_cached_embeddings_for_unchanged_text(corpus: Path) -> None:
    class CountingEmbedder(FakeEmbedder):
        def __init__(self) -> None:
            self.embedded: List[str] = []

        def embed_documents(self, texts, on_batch=None):
            self.embedded.extend(texts)
            return super().embed_documents(texts, on_batch)

    first = CorpusIndex.build(load_corpus(corpus), FakeEmbedder())
    (corpus / "documents" / "DOC-003_Synthetic.txt").write_text("A registered power of attorney is attached to this bid.", encoding="utf-8")
    counting = CountingEmbedder()

    second = CorpusIndex.build(load_corpus(corpus), counting, cache=first.embedding_cache("fake-embed"))

    assert len(counting.embedded) == 1 and "power of attorney" in counting.embedded[0]
    assert second.meta["embeddings_reused"] == len(second.chunks) - 1
    assert first.embedding_cache("another-model") == {}
