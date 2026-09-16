"""Tests for the generated reports.

These reports are the deliverable a supervisor or a procurement officer reads.
What is asserted here is that they follow the team's Document Format Standard
(Times New Roman 12, 1.5 spacing, justified, numbered Heading styles, each main
section on a new page, captions above tables, no page number on the first
page), that every Not Found item says what happened, and that the PDF and the
Word copy carry the same content.
"""

from __future__ import annotations

import unicodedata
from pathlib import Path

import pytest

from procurecheck import checklists
from procurecheck.models import ClauseVerification, CompletenessReport
from procurecheck.report import to_csv, to_dict, to_text
from procurecheck.report_content import CheckMeta, EvaluationMeta, check_report_blocks
from procurecheck.report_docx import build_evaluation_document, write_check_docx, write_evaluation_docx
from procurecheck.report_layout import Heading, Table, build_word
from procurecheck.report_pdf import write_check_pdf, write_evaluation_pdf


class FakeResult:
    """Stands in for run_evaluation.CaseResult, which lives outside the package."""

    def __init__(self, identifier: str, passed: bool):
        self.id = identifier
        self.acceptance_criteria = "AC4"
        self.submission = "synthetic-submission.pdf"
        self.description = "Item deliberately omitted from the submission."
        self.expected = "Not Found, with no page number"
        self.actual = "Found, confidence 0.95. The model was given all 7 pages.\nCHK-05: Found."
        self.observation = "Expected Not Found. The system returned Found."
        self.passed = passed
        self.seconds = 12.0


@pytest.fixture
def report() -> CompletenessReport:
    return CompletenessReport(
        submission_id="synthetic-submission.pdf",
        verified_items=[
            ClauseVerification(
                checklist_item_id="STD-02",
                clause_title="Certificate of incorporation",
                is_present=True,
                page_number=2,
                extracted_snippet="Appendix A is the certificate of incorporation.",
                confidence_score=0.95,
                requires_human_review=False,
            ),
            ClauseVerification(
                checklist_item_id="STD-09",
                clause_title="Bid security",
                is_present=False,
                confidence_score=0.0,
                requires_human_review=False,
            ),
        ],
        missing_items=["STD-09"],
    )


@pytest.fixture
def meta() -> CheckMeta:
    return CheckMeta(
        submission_name="synthetic-submission.pdf",
        model="llama3.1:8b",
        page_count=7,
        checklist_label=checklists.label_for(checklists.STANDARD),
        caveat=checklists.TEMPLATE_CAVEAT,
        pipeline_label="v2.0-per-item + adjudicator-v1.0",
        threshold=0.85,
        evidence_check=True,
    )


@pytest.fixture
def results():
    return [FakeResult("EV-01", True), FakeResult("EV-05", False)]


@pytest.fixture
def evaluation_meta() -> EvaluationMeta:
    return EvaluationMeta(
        model="llama3.1:8b",
        prompt_version="v2.0-per-item",
        strategy="one check per checklist item",
        context_tokens=16384,
        threshold=0.85,
        submission_note="A seven-page synthetic tender submission.",
    )


def _pdf_text(path: Path) -> list[str]:
    pymupdf = pytest.importorskip("pymupdf")
    with pymupdf.open(path) as document:
        # NFKC folds the fi and ffi ligatures back into letters, as a viewer's
        # search does. A space or hyphen that does not survive is a real fault.
        return [unicodedata.normalize("NFKC", page.get_text()) for page in document]


# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------


def test_terminal_report_puts_one_item_on_one_line(report):
    lines = [line for line in to_text(report, "llama3.1:8b").splitlines() if line]

    assert lines[0].startswith("synthetic-submission.pdf")
    assert any(line.startswith("STD-02") and "Found" in line for line in lines)
    assert any(line.startswith("STD-09") and "Not Found" in line for line in lines)


def test_csv_and_json_say_what_happened_to_every_item(report):
    assert "what_happened" in to_csv(report, "llama3.1:8b").splitlines()[5]
    item = to_dict(report, "llama3.1:8b")["items"][1]
    assert "no page and no quotation" in item["what_happened"]
    assert "Search the submission by hand" in item["next_step"]


def test_every_item_explains_what_had_to_be_present_and_what_happened(report, meta):
    texts = [getattr(b, "text", "") for b in check_report_blocks(report, meta)]
    labels = [getattr(b, "label", "") for b in check_report_blocks(report, meta)]

    assert labels.count("What had to be present.") == 2
    assert labels.count("What the system did.") == 2
    assert labels.count("Next step.") == 2
    assert any("all 7 pages" in t and "no page and no quotation" in t for t in texts)


def test_every_table_is_referred_to_before_it_appears(report, meta):
    blocks = check_report_blocks(report, meta)
    for index, block in enumerate(blocks):
        if isinstance(block, Table):
            earlier = " ".join(getattr(b, "text", "") for b in blocks[:index])
            assert f"Table {block.number}" in earlier


def test_sections_and_subsections_are_numbered_in_order(report, meta):
    headings = [b.text for b in check_report_blocks(report, meta) if isinstance(b, Heading)]

    assert headings[0] == "1. Introduction"
    assert "2. Summary of findings" in headings
    assert "3.1 STD-02: Certificate of incorporation" in headings
    assert "3.2 STD-09: Bid security" in headings


def test_a_report_of_almost_nothing_found_asks_for_the_inputs_to_be_checked(meta):
    absent = [
        ClauseVerification(
            checklist_item_id=f"STD-{n:02d}", clause_title=f"Document {n}",
            is_present=False, confidence_score=0.0, requires_human_review=False,
        )
        for n in range(1, 11)
    ]
    report = CompletenessReport(submission_id="x", verified_items=absent, missing_items=[])
    labels = [getattr(b, "label", "") for b in check_report_blocks(report, meta)]

    assert "Check the inputs first." in labels


def test_no_report_text_contains_an_em_or_en_dash(report, meta):
    for block in check_report_blocks(report, meta):
        for value in vars(block).values():
            assert "\u2014" not in str(value) and "\u2013" not in str(value)


# ---------------------------------------------------------------------------
# Format standard, Word
# ---------------------------------------------------------------------------


def test_word_copy_follows_the_document_format_standard(report, meta):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    document = build_word(check_report_blocks(report, meta))
    normal = document.styles["Normal"]
    heading = document.styles["Heading 1"]

    assert normal.font.name == "Times New Roman"
    assert normal.font.size.pt == 12
    assert normal.paragraph_format.line_spacing == 1.5
    assert normal.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert heading.font.name == "Times New Roman" and heading.font.bold
    assert heading.paragraph_format.page_break_before, "each main section starts a new page"
    assert document.sections[0].different_first_page_header_footer, "no number on page one"
    assert "PAGE" in document.sections[0].footer._element.xml


def test_word_captions_sit_above_their_tables(report, meta):
    document = build_word(check_report_blocks(report, meta))
    body = list(document.element.body.iterchildren())
    for index, element in enumerate(body):
        if element.tag.endswith("}tbl"):
            caption = body[index - 1]
            text = "".join(caption.itertext())
            assert text.startswith("Table ") and ":" in text


def test_word_headings_use_real_heading_styles(report, meta):
    document = build_word(check_report_blocks(report, meta))
    styled = {p.style.name for p in document.paragraphs if p.text.startswith(("1. ", "3.1 "))}
    assert styled == {"Heading 1", "Heading 2"}


def test_write_check_docx_produces_a_readable_file(report, meta, tmp_path: Path):
    docx = pytest.importorskip("docx")
    destination = write_check_docx(report, meta, tmp_path / "check.docx")

    reopened = docx.Document(str(destination))
    assert any("What the system did." in p.text for p in reopened.paragraphs)


# ---------------------------------------------------------------------------
# Format standard, PDF
# ---------------------------------------------------------------------------


def test_check_pdf_names_the_checklist_and_carries_the_template_caveat(report, meta, tmp_path: Path):
    pages = _pdf_text(write_check_pdf(report, meta, tmp_path / "check.pdf"))
    text = "".join(pages)

    assert "Completeness Check Report" in pages[0]
    assert "bundled template" in text
    assert "bidding document" in text
    assert "What the system did." in text


def test_check_pdf_omits_the_caveat_when_a_real_checklist_was_used(report, tmp_path: Path):
    meta = CheckMeta(submission_name="bid.pdf", model="llama3.1:8b", page_count=7, checklist_label="tender-42.csv")
    text = "".join(_pdf_text(write_check_pdf(report, meta, tmp_path / "check.pdf")))

    assert "tender-42.csv" in text
    assert "Edit the checklist" not in text


def test_pdf_title_page_is_unnumbered_and_sections_start_new_pages(report, meta, tmp_path: Path):
    pages = _pdf_text(write_check_pdf(report, meta, tmp_path / "check.pdf"))

    assert pages[0].strip().splitlines()[-1] != "1"
    assert pages[1].strip().splitlines()[-1] == "2"
    assert "\xa0" not in pages[1] and "\xad" not in pages[1], "the text layer must stay searchable"
    starts = [page.lstrip().splitlines()[0] for page in pages[1:]]
    for heading in ("1. Introduction", "2. Summary of findings", "3. Item-by-item findings"):
        assert heading in starts


# ---------------------------------------------------------------------------
# Evaluation report
# ---------------------------------------------------------------------------


def test_evaluation_report_sets_expected_beside_actual_with_an_observation(results, evaluation_meta):
    document = build_evaluation_document(results, evaluation_meta, ["One caveat."])
    texts = [p.text for p in document.paragraphs]

    assert "2. Methodology" in texts
    assert "4.2 EV-05: Noticing documents that are missing" in texts
    assert any(t.startswith("Expected behaviour.") for t in texts)
    assert any(t.startswith("Actual behaviour.") for t in texts)
    assert any(t.startswith("Observation.") for t in texts)
    assert "CHK-05: Found." in texts, "multi-line actual behaviour becomes a list"


def test_evaluation_report_discusses_failures_only_when_there_are_any(results, evaluation_meta):
    failing = [p.text for p in build_evaluation_document(results, evaluation_meta, []).paragraphs]
    passing = [p.text for p in build_evaluation_document([FakeResult("EV-01", True)], evaluation_meta, []).paragraphs]

    assert "5. Discussion of failures" in failing
    assert not any(t.endswith("Discussion of failures") for t in passing)
    assert "5. Limitations" in passing


def test_evaluation_pdf_and_word_copy_are_both_written(results, evaluation_meta, tmp_path: Path):
    pdf = write_evaluation_pdf(results, evaluation_meta, ["One caveat."], tmp_path / "e.pdf")
    docx = write_evaluation_docx(results, evaluation_meta, ["One caveat."], tmp_path / "e.docx")

    assert pdf.is_file() and docx.is_file()
    assert "1 of 2 cases behaved as expected" in "".join(_pdf_text(pdf))
