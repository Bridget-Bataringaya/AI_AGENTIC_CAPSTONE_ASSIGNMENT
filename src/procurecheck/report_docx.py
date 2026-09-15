"""Render an evaluation run as a Word document in the team's academic format.

The Markdown table is the cheap, diff-able record kept in the repository. This
is the copy a supervisor or marker opens: same content, same order, same
brevity, but with a title block, numbered sections, captioned tables and page
numbers, which is what the course expects a submitted document to look like.

Every deliverable ships as both, and neither is generated without the other.

Built with python-docx, already a project dependency.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(12)
TABLE_SIZE = Pt(9)
HEADING_COLOUR = RGBColor(0x0F, 0x2B, 0x46)
PASS_COLOUR = RGBColor(0x1E, 0x6B, 0x2B)
FAIL_COLOUR = RGBColor(0xA1, 0x20, 0x20)

INSTITUTION = "Makerere University"
COURSE = "BSE4104 AI-Native and Agentic Engineering Capstone"
GROUP = "Group-H (Evening)"
PROJECT = "Public Procurement Document-Completeness Agent"

# Mirrors report_pdf.GROUP_TITLES so the two documents group cases identically.
GROUP_TITLES: Sequence[Tuple[str, str]] = (
    ("present", "Finding documents that are present"),
    ("absent", "Noticing documents that are missing"),
    ("safety", "Refusing to go beyond its role"),
    ("formats", "Handling different file types"),
    ("injection", "Resisting hidden instructions"),
)


@dataclass(frozen=True)
class DocxMeta:
    """What the run was, stated once at the top of the document."""

    model: str
    prompt_version: str
    strategy: str
    context_tokens: int
    threshold: float
    submission_note: str


def _now() -> str:
    return datetime.now(timezone.utc).strftime("%d %B %Y")


def _flatten(text: object) -> str:
    return " ".join(str(text).split())


def _set_base_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    paragraph_format = normal.paragraph_format
    paragraph_format.space_after = Pt(6)
    paragraph_format.line_spacing = 1.15

    for level in (1, 2):
        style = document.styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style.font.color.rgb = HEADING_COLOUR
        style.font.size = Pt(14 if level == 1 else 12)
        style.font.bold = True

    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def _add_page_number_footer(document: Document) -> None:
    """Put 'Page N' in the footer using a real Word field, not a fixed number."""
    for section in document.sections:
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("Page ")
        run.font.size = Pt(9)

        field_run = paragraph.add_run()
        field_run.font.size = Pt(9)
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = "PAGE"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        for element in (begin, instruction, end):
            field_run._r.append(element)


def _title_block(document: Document, title: str, meta: DocxMeta) -> None:
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = HEADING_COLOUR

    for text, size, bold in (
        (PROJECT, 13, True),
        (COURSE, 11, False),
        (f"{GROUP}, {INSTITUTION}", 11, False),
        (_now(), 10, False),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        line = paragraph.add_run(text)
        line.bold = bold
        line.font.size = Pt(size)

    document.add_paragraph()


def _style_header_row(table) -> None:
    for cell in table.rows[0].cells:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E8EDF2")
        cell._tc.get_or_add_tcPr().append(shading)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True


def _new_table(document: Document, headers: Sequence[str]) -> object:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    _shrink(table)
    _style_header_row(table)
    return table


def _shrink(table, size: Pt = TABLE_SIZE) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = size


def _caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def _fill_row(row, values: Sequence[str]) -> None:
    for cell, value in zip(row.cells, values):
        cell.text = str(value)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(2)
            for run in paragraph.runs:
                run.font.size = TABLE_SIZE


def _verdict_row(row, index: int, passed: bool) -> None:
    """Colour just the verdict cell, so failures are findable at a glance."""
    for paragraph in row.cells[index].paragraphs:
        for run in paragraph.runs:
            run.bold = True
            run.font.color.rgb = PASS_COLOUR if passed else FAIL_COLOUR


def _group_results(results: Sequence) -> Dict[str, List]:
    buckets: Dict[str, List] = {key: [] for key, _ in GROUP_TITLES}
    for result in results:
        if result.id in ("EV-01", "EV-02", "EV-03", "EV-04"):
            buckets["present"].append(result)
        elif result.id in ("EV-05", "EV-06", "EV-07", "EV-18"):
            buckets["absent"].append(result)
        elif result.id in ("EV-08", "EV-09", "EV-10", "EV-11", "EV-13"):
            buckets["safety"].append(result)
        elif result.id == "EV-14":
            buckets["injection"].append(result)
        else:
            buckets["formats"].append(result)
    return buckets


def build_document(
    results: Sequence,
    meta: DocxMeta,
    caveats: Sequence[str],
    title: str = "Evaluation Results",
) -> Document:
    document = Document()
    _set_base_styles(document)
    _add_page_number_footer(document)
    _title_block(document, title, meta)

    passed = sum(1 for result in results if result.passed)
    summary = document.add_paragraph()
    summary_run = summary.add_run(
        f"{passed} of {len(results)} cases behaved as expected."
    )
    summary_run.bold = True
    summary_run.font.size = Pt(13)

    document.add_heading("1. Run details", level=1)
    details = _new_table(document, ["Setting", "Value"])
    for label, value in (
        ("Model", meta.model),
        ("Prompt version", meta.prompt_version),
        ("Matching method", meta.strategy),
        ("Human-review threshold", f"{meta.threshold}"),
        ("Context window", f"{meta.context_tokens:,} tokens"),
        ("Cases run", f"{len(results)}"),
        ("Cases meeting expectation", f"{passed} of {len(results)}"),
        ("Documents tested", _flatten(meta.submission_note)),
    ):
        _fill_row(details.add_row(), (label, value))
    _caption(document, "Table 1: Configuration of the evaluated run.")

    document.add_heading("2. Results by area", level=1)
    area = _new_table(document, ["Area tested", "As expected"])
    groups = _group_results(results)
    for key, group_title in GROUP_TITLES:
        members = groups.get(key, [])
        if not members:
            continue
        group_passed = sum(1 for result in members if result.passed)
        row = area.add_row()
        _fill_row(row, (group_title, f"{group_passed} of {len(members)}"))
        _verdict_row(row, 1, group_passed == len(members))
    _caption(document, "Table 2: Cases meeting expectation, grouped by area.")

    failures = [result for result in results if not result.passed]
    next_section = 3
    if failures:
        document.add_heading(f"{next_section}. Failures", level=1)
        for result in failures:
            paragraph = document.add_paragraph(style="List Bullet")
            identifier = paragraph.add_run(f"{result.id}. ")
            identifier.bold = True
            paragraph.add_run(
                f"{_flatten(result.description)} Expected "
                f"{_flatten(result.expected)}. Got {_flatten(result.actual)}."
            )
            for run in paragraph.runs:
                run.font.size = Pt(10)
        next_section += 1

    document.add_heading(f"{next_section}. Every case", level=1)
    cases = _new_table(
        document, ["Case", "AC", "Scenario", "Expected", "Actual", "Result"]
    )
    for result in results:
        row = cases.add_row()
        _fill_row(
            row,
            (
                result.id,
                _flatten(getattr(result, "acceptance_criteria", "")),
                _flatten(result.description),
                _flatten(result.expected),
                _flatten(result.actual),
                "Pass" if result.passed else "Fail",
            ),
        )
        _verdict_row(row, 5, result.passed)
    _caption(document, "Table 3: Expected against actual behaviour, case by case.")
    next_section += 1

    document.add_heading(f"{next_section}. Limitations", level=1)
    for caveat in caveats:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(_flatten(caveat)).font.size = Pt(10)

    return document


def write_docx(
    results: Sequence,
    meta: DocxMeta,
    caveats: Sequence[str],
    destination: Path,
    title: str = "Evaluation Results",
) -> Path:
    """Write the evaluation results to `destination` as a Word document."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    build_document(results, meta, caveats, title).save(str(destination))
    return destination
