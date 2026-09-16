"""One document layout, rendered to Word and to PDF.

Every report this project produces follows the team's Document Format Standard
(docs shared from the SEP I repository, "Document Format Standard.docx"):

- Times New Roman throughout, body text 12 point, line spacing 1.5, justified
- Numbered headings in the word processor's real Heading styles, bold, larger
  than body text
- Each main section starts on a new page
- Table captions above the table, numbered Table 1, Table 2, and every table
  referred to in the text before it appears
- Page numbers in the footer, none on the first page
- The group shown on the document

A report is built once as a list of blocks and then rendered twice. That is
what keeps the PDF, which is the system's output of record, and the Word copy,
which is the editable working document, saying the same thing in the same
order. Numbering is assigned while the blocks are built, so the text that
introduces a table can name its number before the table exists.

Word output uses python-docx, in Times New Roman. PDF output uses PyMuPDF's
Story API, in MuPDF's built-in Times (Nimbus Roman), which has the same design
and metrics. Embedding the Windows Times New Roman file was tried and dropped:
its character map sends every space to a non-breaking space and every hyphen to
a soft hyphen, so the PDF of record could not be searched or copied from. Both
libraries are existing project dependencies.
"""

from __future__ import annotations

import html
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Sequence, Tuple, Union

INSTITUTION = "Makerere University"
COURSE = "BSE4104 AI-Native and Agentic Engineering Capstone"
GROUP = "Group H (Evening)"
PROJECT = "Public Procurement Document-Completeness Agent"

BODY_FONT = "Times New Roman"
LINE_SPACING = 1.5

# Tones for a table cell whose value is a verdict. Colour supports the word, it
# never replaces it: every coloured cell also says Pass, Fail, Found and so on.
TONE_GOOD = "good"
TONE_BAD = "bad"
TONE_WARN = "warn"
TONE_RGB = {TONE_GOOD: (0x1E, 0x6B, 0x2B), TONE_BAD: (0xA1, 0x20, 0x20), TONE_WARN: (0x8A, 0x5A, 0x00)}


# ---------------------------------------------------------------------------
# Blocks
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Cell:
    text: str
    tone: Optional[str] = None
    bold: bool = False


CellValue = Union[str, Cell]


@dataclass(frozen=True)
class TitlePage:
    title: str
    subtitle: str
    lines: Tuple[str, ...]


@dataclass(frozen=True)
class Heading:
    text: str
    level: int


@dataclass(frozen=True)
class Paragraph:
    text: str
    label: str = ""
    italic: bool = False
    indent: bool = False


@dataclass(frozen=True)
class BulletList:
    items: Tuple[str, ...]


@dataclass(frozen=True)
class Table:
    number: int
    caption: str
    headers: Tuple[str, ...]
    rows: Tuple[Tuple[CellValue, ...], ...]
    # Relative column widths. Word needs them; the PDF uses them as percentages.
    widths: Tuple[float, ...] = ()


Block = Union[TitlePage, Heading, Paragraph, BulletList, Table]


@dataclass
class ReportBuilder:
    """Accumulates blocks and hands out section and table numbers in order."""

    blocks: List[Block] = field(default_factory=list)
    _section: int = 0
    _subsection: int = 0
    _tables: int = 0

    def title_page(self, title: str, subtitle: str, date_line: str) -> None:
        self.blocks.append(
            TitlePage(
                title=title,
                subtitle=subtitle,
                lines=(PROJECT, COURSE, f"{GROUP}, {INSTITUTION}", date_line),
            )
        )

    def section(self, text: str) -> None:
        self._section += 1
        self._subsection = 0
        self.blocks.append(Heading(f"{self._section}. {text}", 1))

    def subsection(self, text: str) -> None:
        self._subsection += 1
        self.blocks.append(Heading(f"{self._section}.{self._subsection} {text}", 2))

    def paragraph(self, text: str, label: str = "") -> None:
        self.blocks.append(Paragraph(text=text, label=label))

    def quote(self, text: str) -> None:
        self.blocks.append(Paragraph(text=text, italic=True, indent=True))

    def bullets(self, items: Sequence[str]) -> None:
        if items:
            self.blocks.append(BulletList(tuple(items)))

    @property
    def next_table(self) -> int:
        """The number the next table will carry, for the sentence introducing it."""
        return self._tables + 1

    def table(
        self,
        caption: str,
        headers: Sequence[str],
        rows: Sequence[Sequence[CellValue]],
        widths: Sequence[float] = (),
    ) -> int:
        self._tables += 1
        self.blocks.append(
            Table(
                number=self._tables,
                caption=caption,
                headers=tuple(headers),
                rows=tuple(tuple(row) for row in rows),
                widths=tuple(widths),
            )
        )
        return self._tables


def _cell(value: CellValue) -> Cell:
    return value if isinstance(value, Cell) else Cell(str(value))


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------


def _strip_theme_fonts(style) -> None:
    """Remove theme font references, which override an explicit font name.

    Word's built-in heading styles point at the theme's heading font. Setting
    font.name adds Times New Roman beside that reference rather than replacing
    it, and Word honours the theme, so headings silently render in Calibri
    Light unless the reference is removed.
    """
    from docx.oxml.ns import qn

    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        return
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attribute), None)
    fonts.set(qn("w:eastAsia"), BODY_FONT)
    fonts.set(qn("w:cs"), BODY_FONT)


def _add_field(paragraph, instruction: str, cached: str = ""):
    """Insert a Word field such as PAGE or SEQ, with a result shown before update."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def char(kind: str):
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        return element

    begin_run = paragraph.add_run()
    begin_run._r.append(char("begin"))
    instruction_run = paragraph.add_run()
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = f" {instruction} "
    instruction_run._r.append(text)
    paragraph.add_run()._r.append(char("separate"))
    result_run = paragraph.add_run(cached)
    paragraph.add_run()._r.append(char("end"))
    return result_run


def _configure_word(document) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt, RGBColor

    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = section.right_margin = Mm(25.4)
        section.top_margin = section.bottom_margin = Mm(25.4)
        # The first page carries no number, per the standard.
        section.different_first_page_header_footer = True
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        number = _add_field(footer, "PAGE", "2")
        number.font.name = BODY_FONT
        number.font.size = Pt(12)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(12)
    _strip_theme_fonts(normal)
    normal.paragraph_format.line_spacing = LINE_SPACING
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for level, size in ((1, 16), (2, 14)):
        style = document.styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        _strip_theme_fonts(style)
        style.paragraph_format.line_spacing = LINE_SPACING
        style.paragraph_format.space_before = Pt(12 if level == 2 else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # Each main section starts on a new page.
        style.paragraph_format.page_break_before = level == 1

    caption = document.styles["Caption"]
    caption.font.name = BODY_FONT
    caption.font.size = Pt(12)
    caption.font.italic = False
    caption.font.bold = False
    caption.font.color.rgb = RGBColor(0, 0, 0)
    _strip_theme_fonts(caption)
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    bullet = document.styles["List Bullet"]
    bullet.font.name = BODY_FONT
    bullet.font.size = Pt(12)
    bullet.paragraph_format.line_spacing = LINE_SPACING
    bullet.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _word_title_page(document, block: TitlePage) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    for _ in range(5):
        document.add_paragraph()
    for text, size, bold in (
        (block.title, 20, True),
        (block.subtitle, 14, False),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
    document.add_paragraph()
    for index, text in enumerate(block.lines):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = index == 0


def _word_table(document, block: Table) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    caption = document.add_paragraph(style="Caption")
    label = caption.add_run("Table ")
    label.bold = True
    number = _add_field(caption, "SEQ Table \\* ARABIC", str(block.number))
    number.bold = True
    colon = caption.add_run(": ")
    colon.bold = True
    caption.add_run(block.caption)

    table = document.add_table(rows=1, cols=len(block.headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    text_width_mm = 210 - 2 * 25.4
    widths = block.widths or tuple(1.0 for _ in block.headers)
    total = sum(widths)
    column_widths = [Mm(text_width_mm * w / total) for w in widths]

    def fill(cell, value: Cell, header: bool, width) -> None:
        cell.width = width
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(value.text)
        run.font.name = BODY_FONT
        run.font.size = Pt(10)
        run.bold = header or value.bold or value.tone is not None
        if value.tone in TONE_RGB:
            run.font.color.rgb = RGBColor(*TONE_RGB[value.tone])

    header_row = table.rows[0]
    # Repeat the header row when the table runs onto another page.
    header_properties = header_row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_properties.append(repeat)
    for cell, header, width in zip(header_row.cells, block.headers, column_widths):
        fill(cell, Cell(header), True, width)

    for row in block.rows:
        cells = table.add_row().cells
        for cell, value, width in zip(cells, row, column_widths):
            fill(cell, _cell(value), False, width)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def build_word(blocks: Sequence[Block]):
    """Render blocks into a python-docx Document."""
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Mm

    document = Document()
    _configure_word(document)

    for block in blocks:
        if isinstance(block, TitlePage):
            _word_title_page(document, block)
            # Only needed when no numbered section follows, since a Heading 1
            # already starts a new page.
            if not any(isinstance(b, Heading) and b.level == 1 for b in blocks):
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif isinstance(block, Heading):
            document.add_heading(block.text, level=block.level)
        elif isinstance(block, Paragraph):
            paragraph = document.add_paragraph()
            if block.indent:
                paragraph.paragraph_format.left_indent = Mm(12.7)
            if block.label:
                paragraph.add_run(block.label + " ").bold = True
            run = paragraph.add_run(block.text)
            run.italic = block.italic
        elif isinstance(block, BulletList):
            for item in block.items:
                document.add_paragraph(item, style="List Bullet")
        elif isinstance(block, Table):
            _word_table(document, block)
    return document


def write_word(blocks: Sequence[Block], destination: Path) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    build_word(blocks).save(str(destination))
    return destination


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

A4_WIDTH = 595
A4_HEIGHT = 842
PDF_MARGIN = 72  # one inch, matching the Word copy

PDF_CSS = """
body { font-family: Times, serif; font-size: 12pt; line-height: 1.5; color: #1a1a1a; text-align: justify; }
h1 { font-size: 16pt; font-weight: bold; margin: 0 0 8pt 0; text-align: left; line-height: 1.3; }
h2 { font-size: 14pt; font-weight: bold; margin: 12pt 0 6pt 0; text-align: left; line-height: 1.3; }
p { margin: 0 0 6pt 0; }
.quote { margin-left: 36pt; font-style: italic; }
.caption { margin: 6pt 0 4pt 0; text-align: left; }
ul { margin: 0 0 6pt 0; }
li { margin: 0 0 3pt 0; }
table { width: 100%; border-collapse: collapse; margin: 0 0 10pt 0; }
th { font-size: 10pt; line-height: 1.15; text-align: left; vertical-align: top;
     font-weight: bold; border: 1px solid #555555; padding: 3pt; }
td { font-size: 10pt; line-height: 1.15; text-align: left; vertical-align: top;
     border: 1px solid #555555; padding: 3pt; }
.good { color: #1e6b2b; font-weight: bold; }
.bad { color: #a12020; font-weight: bold; }
.warn { color: #8a5a00; font-weight: bold; }
.strong { font-weight: bold; }
.title { text-align: center; }
.title-main { font-size: 20pt; font-weight: bold; margin: 190pt 0 10pt 0; text-align: center; line-height: 1.3; }
.title-sub { font-size: 14pt; margin: 0 0 30pt 0; text-align: center; line-height: 1.3; }
.title-line { text-align: center; margin: 0 0 4pt 0; }
"""


def _e(text: object) -> str:
    return html.escape(str(text), quote=False)


def _html_block(block: Block) -> str:
    if isinstance(block, TitlePage):
        lines = "".join(
            f'<p class="title-line">{"<b>" if i == 0 else ""}{_e(line)}{"</b>" if i == 0 else ""}</p>'
            for i, line in enumerate(block.lines)
        )
        return (
            f'<p class="title-main">{_e(block.title)}</p>'
            f'<p class="title-sub">{_e(block.subtitle)}</p>{lines}'
        )
    if isinstance(block, Heading):
        tag = "h1" if block.level == 1 else "h2"
        return f"<{tag}>{_e(block.text)}</{tag}>"
    if isinstance(block, Paragraph):
        css = ' class="quote"' if block.indent or block.italic else ""
        label = f"<b>{_e(block.label)}</b> " if block.label else ""
        return f"<p{css}>{label}{_e(block.text)}</p>"
    if isinstance(block, BulletList):
        return "<ul>" + "".join(f"<li>{_e(item)}</li>" for item in block.items) + "</ul>"
    if isinstance(block, Table):
        widths = block.widths or tuple(1.0 for _ in block.headers)
        total = sum(widths)
        head = "".join(
            f'<th width="{100 * w / total:.0f}%">{_e(h)}</th>'
            for h, w in zip(block.headers, widths)
        )
        body = []
        for row in block.rows:
            cells = []
            for value in row:
                cell = _cell(value)
                css = cell.tone or ("strong" if cell.bold else "")
                content = f'<span class="{css}">{_e(cell.text)}</span>' if css else _e(cell.text)
                cells.append(f"<td>{content}</td>")
            body.append("<tr>" + "".join(cells) + "</tr>")
        return (
            f'<p class="caption"><b>Table {block.number}:</b> {_e(block.caption)}</p>'
            f"<table><tr>{head}</tr>{''.join(body)}</table>"
        )
    raise TypeError(f"unknown block {block!r}")


def _chunks(blocks: Sequence[Block]) -> List[str]:
    """Split the document so the title page and every main section start a page."""
    chunks: List[List[Block]] = []
    for block in blocks:
        starts_page = isinstance(block, TitlePage) or (
            isinstance(block, Heading) and block.level == 1
        )
        if starts_page or not chunks:
            chunks.append([])
        chunks[-1].append(block)
    return ["".join(_html_block(b) for b in chunk) for chunk in chunks]


def _pymupdf():
    try:
        import pymupdf
    except ImportError:  # PyMuPDF older than 1.24 only exposes `fitz`
        import fitz as pymupdf
    return pymupdf


def write_pdf(blocks: Sequence[Block], destination: Path) -> Path:
    """Render blocks to an A4 PDF, numbering every page except the first."""
    pymupdf = _pymupdf()
    destination.parent.mkdir(parents=True, exist_ok=True)
    frame = pymupdf.Rect(PDF_MARGIN, PDF_MARGIN, A4_WIDTH - PDF_MARGIN, A4_HEIGHT - PDF_MARGIN)
    mediabox = pymupdf.Rect(0, 0, A4_WIDTH, A4_HEIGHT)

    writer = pymupdf.DocumentWriter(str(destination))
    for chunk in _chunks(blocks):
        story = pymupdf.Story(html=chunk, user_css=PDF_CSS)
        more = True
        while more:
            device = writer.begin_page(mediabox)
            more, _ = story.place(frame)
            story.draw(device)
            writer.end_page()
    writer.close()

    document = pymupdf.open(destination)
    for index, page in enumerate(document, start=1):
        if index == 1:
            continue
        label = str(index)
        width = pymupdf.get_text_length(label, fontname="tiro", fontsize=12)
        page.insert_text(
            ((A4_WIDTH - width) / 2, A4_HEIGHT - 36),
            label,
            fontname="tiro",
            fontsize=12,
        )
    document.saveIncr()
    document.close()
    return destination
