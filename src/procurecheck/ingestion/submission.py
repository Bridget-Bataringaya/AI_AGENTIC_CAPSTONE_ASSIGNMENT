"""Submission Ingestion and Parsing.

Extracts text from a tender submission and preserves page markers in the form
"[PAGE n]" so that the matching engine can report an accurate page_number, as
required by User Stories AC3 and Prompt Specification v1.0 Sec. 4.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Final, List, Sequence, Tuple

PAGE_MARKER_TEMPLATE: Final[str] = "[PAGE {number}]"
SUPPORTED_SUFFIXES: Final[Tuple[str, ...]] = (".pdf", ".txt", ".md", ".docx")

# A page yielding fewer characters than this is very likely a scanned image
# rather than a text layer, and needs OCR before it can be matched.
MIN_CHARS_FOR_TEXT_PAGE: Final[int] = 20


class UnsupportedDocumentError(ValueError):
    """Raised when a file type cannot be parsed."""


class EmptyDocumentError(ValueError):
    """Raised when a file parses successfully but yields no usable text."""


@dataclass(frozen=True)
class ParsedPage:
    number: int
    text: str

    @property
    def looks_scanned(self) -> bool:
        return len(self.text.strip()) < MIN_CHARS_FOR_TEXT_PAGE


@dataclass(frozen=True)
class ParsedSubmission:
    """A parsed submission ready to be passed to the matching engine."""

    submission_id: str
    pages: Sequence[ParsedPage]

    @property
    def page_count(self) -> int:
        return len(self.pages)

    @property
    def scanned_pages(self) -> List[int]:
        """Pages with too little text to match reliably. Surfaced to the user
        rather than silently treated as empty, per the Charter's
        document-quality constraint."""
        return [page.number for page in self.pages if page.looks_scanned]

    @property
    def character_count(self) -> int:
        return sum(len(page.text) for page in self.pages)

    def as_marked_text(self) -> str:
        """Render the whole submission with page markers preserved."""
        return "\n\n".join(
            f"{PAGE_MARKER_TEMPLATE.format(number=page.number)}\n{page.text.strip()}"
            for page in self.pages
        )


def _read_pdf(path: Path) -> List[ParsedPage]:
    import fitz  # PyMuPDF

    pages: List[ParsedPage] = []
    with fitz.open(path) as document:
        for index, page in enumerate(document, start=1):
            pages.append(ParsedPage(number=index, text=page.get_text("text")))
    return pages


def _read_docx(path: Path) -> List[ParsedPage]:
    """Read a .docx file.

    Word documents carry no reliable page boundaries without rendering, so the
    whole document is reported as page 1 rather than inventing page numbers
    that a human reviewer could not verify against the original.
    """
    import docx

    document = docx.Document(str(path))
    blocks: List[str] = [
        paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
    ]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return [ParsedPage(number=1, text="\n".join(blocks))]


def _read_plain_text(path: Path) -> List[ParsedPage]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return [ParsedPage(number=1, text=text)]


def parse_submission(path: Path | str, submission_id: str | None = None) -> ParsedSubmission:
    """Parse a submission file into pages of text.

    Raises UnsupportedDocumentError for unknown file types and
    EmptyDocumentError when nothing usable could be extracted, so that the API
    can return the clear error message required by User Stories AC2 instead of
    proceeding with an empty document.
    """
    path = Path(path)
    if not path.is_file():
        raise UnsupportedDocumentError(f"No such file: {path}")

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedDocumentError(
            f"Unsupported file type {suffix!r}. Supported types: "
            f"{', '.join(SUPPORTED_SUFFIXES)}."
        )

    try:
        if suffix == ".pdf":
            pages = _read_pdf(path)
        elif suffix == ".docx":
            pages = _read_docx(path)
        else:
            pages = _read_plain_text(path)
    except (UnsupportedDocumentError, EmptyDocumentError):
        raise
    except Exception as exc:
        raise UnsupportedDocumentError(
            f"Could not read {path.name}. The file may be corrupted or password "
            f"protected. Underlying error: {exc}"
        ) from exc

    if not any(page.text.strip() for page in pages):
        raise EmptyDocumentError(
            f"{path.name} contains no extractable text. If it is a scanned "
            f"document, it must be passed through OCR before it can be checked."
        )

    return ParsedSubmission(
        submission_id=submission_id or path.stem, pages=tuple(pages)
    )
