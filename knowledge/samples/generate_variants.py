"""Generate the additional synthetic submissions used for coverage testing.

The base seven-page PDF from generate_submission.py exercises only one format
and only one content shape. These variants close the gaps that a single
submission hides:

    synthetic-submission.docx   same content, Word format. .docx carries no
                                reliable page boundaries, so this checks that
                                page numbers degrade honestly to 1 rather than
                                being invented.
    synthetic-submission.txt    same content, plain text.
    scanned-submission.pdf      image-only pages with no text layer, the case
                                the Charter names as a document-quality risk
                                and the architecture doc promises OCR for.
    no-items-submission.pdf     a real-looking tender package containing NONE
                                of the ten required documents. Isolates whether
                                the model can ever say Not Found, or only
                                latches onto the nearest plausible text.
    long-submission.pdf         large enough to exceed the configured context
                                budget, so the refusal path is exercised rather
                                than assumed.

All content is synthetic: no real bidder, no real procuring entity, no
confidential information, per the Project Charter's data-access constraint.

Run:  python knowledge/samples/generate_variants.py
"""

from __future__ import annotations

from pathlib import Path
from typing import List, Tuple

from generate_submission import PAGES

HERE = Path(__file__).parent

# Enough dense pages to pass the default context budget of
# 16384 - 2000 overhead - 2048 output = 12336 tokens, roughly 49,000 characters.
LONG_PAGE_COUNT = 60

# A tender package that is genuinely missing every required document. It reads
# like a real submission so that a Not Found verdict has to come from absence,
# not from the document looking obviously empty.
NO_ITEMS_PAGES: List[Tuple[str, str]] = [
    (
        "SECTION 1 - TECHNICAL METHOD STATEMENT",
        "Tender Reference: SYN/WORKS/2026/099\n"
        "Bidder: Nsereko Engineering Services (fictional entity)\n\n"
        "The works will be executed in four phases. Site clearance and setting\n"
        "out will be completed in week one. Excavation to formation level will\n"
        "follow, with spoil removed to the approved tip. Blinding and\n"
        "reinforcement will be inspected before any concrete is poured.\n"
        "Curing will be maintained for seven days under damp hessian.",
    ),
    (
        "SECTION 2 - PROPOSED CONSTRUCTION PROGRAMME",
        "Phase 1  Site establishment and clearance      Weeks 1 to 2\n"
        "Phase 2  Substructure and foundations          Weeks 3 to 9\n"
        "Phase 3  Superstructure and roofing            Weeks 10 to 20\n"
        "Phase 4  Finishes, testing and handover        Weeks 21 to 26\n\n"
        "Float of three weeks has been allowed against the rainy season. The\n"
        "critical path runs through the substructure works.",
    ),
    (
        "SECTION 3 - PLANT AND EQUIPMENT SCHEDULE",
        "  1 excavator, 20 tonne, owned\n"
        "  2 tipper trucks, 10 cubic metres, owned\n"
        "  1 concrete mixer, 500 litre, owned\n"
        "  1 poker vibrator, owned\n"
        "  1 water pump, 3 inch, to be hired\n"
        "  Scaffolding, 400 square metres, to be hired\n\n"
        "All owned plant is serviced quarterly by the manufacturer's agent.",
    ),
    (
        "SECTION 4 - KEY PERSONNEL",
        "Site Agent, twelve years of experience in building works.\n"
        "Site Engineer, eight years of experience, degree in civil engineering.\n"
        "Foreman, fifteen years of experience in reinforced concrete.\n"
        "Safety Officer, part time, five years of experience.\n\n"
        "Curricula vitae are held on file and can be produced on request.",
    ),
    (
        "SECTION 5 - HEALTH AND SAFETY APPROACH",
        "Personal protective equipment will be issued to all site personnel.\n"
        "A daily briefing will be held before work begins. Excavations deeper\n"
        "than one point two metres will be shored and barricaded. A first aid\n"
        "box and a trained first aider will be present on site at all times.\n"
        "Accidents will be recorded in the site register and reported to the\n"
        "supervising officer within twenty four hours.",
    ),
]


def _pymupdf():
    try:
        import pymupdf
    except ImportError:  # PyMuPDF older than 1.24 only exposes `fitz`
        import fitz as pymupdf
    return pymupdf


def _plain_text() -> str:
    return "\n\n".join(f"{heading}\n{body}" for heading, body in PAGES)


def _write_pdf(pages: List[Tuple[str, str]], destination: Path) -> Path:
    pymupdf = _pymupdf()
    document = pymupdf.open()
    for heading, body in pages:
        page = document.new_page()
        page.insert_text((60, 70), heading, fontname="helv", fontsize=13)
        page.insert_text((60, 105), body, fontname="helv", fontsize=10)
    document.save(destination)
    document.close()
    return destination


def build_docx() -> Path:
    import docx

    destination = HERE / "synthetic-submission.docx"
    document = docx.Document()
    for heading, body in PAGES:
        document.add_heading(heading, level=2)
        for paragraph in body.split("\n\n"):
            document.add_paragraph(paragraph)
    document.save(str(destination))
    return destination


def build_txt() -> Path:
    destination = HERE / "synthetic-submission.txt"
    destination.write_text(_plain_text(), encoding="utf-8")
    return destination


def build_scanned() -> Path:
    """Render the text to images, then rebuild a PDF with no text layer.

    This is what a photocopied or phone-photographed tender pack looks like to
    a parser: visually readable, programmatically empty.
    """
    pymupdf = _pymupdf()
    destination = HERE / "scanned-submission.pdf"

    source = pymupdf.open(HERE / "synthetic-submission.pdf")
    output = pymupdf.open()
    for page in source:
        # Grayscale at screen resolution, stored as JPEG. A lossless colour
        # render of the same pages costs about 24 MB, which has no place in a
        # git repository; this is a few hundred kilobytes and is just as
        # unreadable to a text parser, which is the whole point of the fixture.
        pixmap = page.get_pixmap(dpi=72, colorspace=pymupdf.csGRAY)
        new_page = output.new_page(width=page.rect.width, height=page.rect.height)
        new_page.insert_image(new_page.rect, stream=pixmap.tobytes("jpeg", jpg_quality=60))
    output.save(destination, deflate=True, garbage=4)
    output.close()
    source.close()
    return destination


def build_no_items() -> Path:
    return _write_pdf(NO_ITEMS_PAGES, HERE / "no-items-submission.pdf")


def build_long() -> Path:
    """Repeat dense filler until the document exceeds the context budget."""
    filler_body = "\n".join(
        "Clause {n}.{m}. The Contractor shall comply with the Specification, "
        "the Drawings and the instructions of the Supervising Officer in all "
        "respects, and shall give reasonable notice before covering up any "
        "work required to be inspected.".format(n=1, m=m)
        for m in range(1, 12)
    )
    pages = [
        (f"SECTION {index} - GENERAL CONDITIONS CONTINUED", filler_body)
        for index in range(1, LONG_PAGE_COUNT + 1)
    ]
    return _write_pdf(pages, HERE / "long-submission.pdf")


def build_all() -> List[Path]:
    if not (HERE / "synthetic-submission.pdf").is_file():
        raise SystemExit(
            "Run generate_submission.py first: scanned-submission.pdf is built "
            "from synthetic-submission.pdf."
        )
    return [
        build_docx(),
        build_txt(),
        build_scanned(),
        build_no_items(),
        build_long(),
    ]


if __name__ == "__main__":
    for path in build_all():
        print(f"Wrote {path.name} ({path.stat().st_size:,} bytes)")
